"""Builds the improved eco-run record sheet (本番用記録表 改良版).

Same idea as the team's sheet (enter the cumulative time at each control line,
get the gap and the speed for the next lap), with lap distances that follow the
real course: the standing start is before the control line and the goal is
before it too, so the first and last laps are shorter (checked on last year's
Super Cub GPX, 2025-10-12).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, color="555555", size=9)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDE3EA")
GREEN_F = PatternFill("solid", fgColor="B7E1B0")
RED_F = PatternFill("solid", fgColor="F4B6B0")
AMBER_F = PatternFill("solid", fgColor="FFE699")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TIME_FMT = '[h]:mm:ss.0'
SEC_FMT = '+0.0"秒";-0.0"秒";0.0"秒"'
KMH_FMT = '0.00'
M_FMT = '#,##0.0'

wb = Workbook()
wb.remove(wb.active)


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, color=f.color, size=f.sz)


def race_sheet(name, laps, limit_min, limit_s, distance):
    ws = wb.create_sheet(name)
    ws["B2"] = name
    ws["B2"].font = TITLE

    # Inputs (yellow fill, blue text).
    inputs = [
        ("全体距離（m）", distance, M_FMT, "競技規則 第2章第1条第1項4（2026年）"),
        ("周回数", laps, "0", "競技規則 第2章第1条第1項4（2026年）"),
        ("制限タイム", f"=TIME(0,{limit_min},{int(limit_s)})+{round(limit_s - int(limit_s), 2)}/86400", TIME_FMT,
         "競技規則 第2章第1条第1項4（2026年）。小数点以下も規則の値"),
        ("余裕（秒）", 30, "0", "制限タイムよりこれだけ早くゴールする。昨年のスーパーカブは約26秒早くゴールしていた（2025-10-12のGPX）。"
                            "元の記録表は16:20（約7秒）"),
        ("スタート→コントロール（m）", 277.8, M_FMT, "ラインの間隔（直線距離）。コースウォークで測った値に書き換える。"
                                         "アプリの設定→ラインに出る「コントロール→スタート」の値。今は地図から拾ったおおよその位置"),
        ("ゴール→コントロール（m）", 89.9, M_FMT, "ラインの間隔（直線距離）。コースウォークで測った値に書き換える。"
                                      "アプリの設定→ラインに出る「ゴール→コントロール」の値。今は地図から拾ったおおよその位置"),
    ]
    ws["K3"] = "入力（黄色のセル）"; ws["K3"].font = BOLD
    for i, (label, value, fmt, note) in enumerate(inputs):
        r = 4 + i
        ws.cell(r, 11, label).font = BLACK
        c = ws.cell(r, 13, value)
        c.font = BLUE; c.fill = YELLOW; c.number_format = fmt; c.border = BOX
        c.comment = Comment(note, "EcoLogger")
        ws.cell(r, 14, note).font = NOTE
    D, N, LIMIT, MARGIN, SC, GC = "$M$4", "$M$5", "$M$6", "$M$7", "$M$8", "$M$9"

    ws["K11"] = "計算（触らない）"; ws["K11"].font = BOLD
    derived = [
        ("全体目標タイム", f"={LIMIT}-{MARGIN}/86400", TIME_FMT, "制限タイム − 余裕"),
        ("1周の長さ（m）", f"=({D}+{SC}+{GC})/{N}", M_FMT, "1周目はスタート→コントロールの分、最終周はゴール→コントロールの分だけ短い"),
        ("計画の平均速度（km/h）", f"={D}/1000/($M$12*24)", KMH_FMT, "全体距離 ÷ 全体目標タイム。1周目の目標速度と同じ"),
    ]
    for i, (label, formula, fmt, note) in enumerate(derived):
        r = 12 + i
        ws.cell(r, 11, label).font = BLACK
        c = ws.cell(r, 13, formula); c.font = BLACK; c.number_format = fmt; c.border = BOX
        ws.cell(r, 14, note).font = NOTE
    TARGET, L = "$M$12", "$M$13"

    heads = ["周回", "周の距離（m）", "目標タイム（累計）", "実走タイム（累計）", "差", "周タイム",
             "一周平均速度", "この周の目標速度", "目標比（km/h）", "次の周の目標速度", "次の周の前半の目標速度"]
    hr = 16
    for j, h in enumerate(heads):
        c = ws.cell(hr, 2 + j, h)
        c.font = BOLD; c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[hr].height = 32
    # Row 17: before the start (lap 0), so lap 1's "next" values come from it.
    r0 = hr + 1
    ws.cell(r0, 2, "スタート前").font = BLACK
    ws.cell(r0, 4, 0).number_format = TIME_FMT
    ws.cell(r0, 5, 0).number_format = TIME_FMT
    ws.cell(r0, 5).font = BLACK
    first = r0 + 1
    rows = 7
    for k in range(1, rows + 1):
        r = first + k - 1
        prev = r - 1
        on = f"$B{r}<={N}"
        ws.cell(r, 2, f'=IF({k}<={N},{k},"")')
        ws.cell(r, 3, f'=IF({on},{L}-IF($B{r}=1,{SC},0)-IF($B{r}={N},{GC},0),"")').number_format = M_FMT
        ws.cell(r, 4, f'=IF({on},{TARGET}*SUM($C${first}:C{r})/{D},"")').number_format = TIME_FMT
        e = ws.cell(r, 5)  # input
        e.number_format = TIME_FMT; e.font = BLUE; e.fill = YELLOW; e.border = BOX
        has = f'AND({on},E{r}<>"")'
        ws.cell(r, 6, f'=IF({has},(D{r}-E{r})*86400,"")').number_format = SEC_FMT
        ws.cell(r, 7, f'=IF({has},E{r}-E{prev},"")').number_format = TIME_FMT
        ws.cell(r, 8, f'=IF({has},C{r}/1000/(G{r}*24),"")').number_format = KMH_FMT
        ws.cell(r, 9, f'=IF(AND({on},E{prev}<>""),C{r}/1000/((D{r}-E{prev})*24),"")').number_format = KMH_FMT
        ws.cell(r, 10, f'=IF(AND({has},I{r}<>""),H{r}-I{r},"")').number_format = '+0.00;-0.00;0.00'
        nxt = r + 1
        ws.cell(r, 11, f'=IF(AND({has},$B{r}<{N}),C{nxt}/1000/((D{nxt}-E{r})*24),"")').number_format = KMH_FMT
        ws.cell(r, 12, f'=IF(AND({has},$B{r}<{N}),C{nxt}/2/1000/(((D{r}+D{nxt})/2-E{r})*24),"")').number_format = KMH_FMT
        for j in range(2, 13):
            ws.cell(r, j).border = BOX
            if j != 5:
                ws.cell(r, j).font = BLACK
    # Lap 1's target before the start.
    ws.cell(r0, 11, f"=C{first}/1000/(D{first}*24)").number_format = KMH_FMT
    ws.cell(r0, 12, f"=C{first}/2/1000/(D{first}/2*24)").number_format = KMH_FMT
    last = first + rows - 1

    # Colours: the lap average against that lap's target (green = on target up to +1 km/h,
    # red = slow, yellow = too fast), and the gap (green = ahead, red = behind).
    rng = f"H{first}:H{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(H{first}<>"",I{first}<>"",H{first}<I{first})'], fill=RED_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(H{first}<>"",I{first}<>"",H{first}>I{first}+1)'], fill=AMBER_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(H{first}<>"",I{first}<>"")'], fill=GREEN_F))
    rng = f"F{first}:F{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(F{first}<>"",F{first}<0)'], fill=RED_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(F{first}<>"",F{first}>=0)'], fill=GREEN_F))

    n = last + 2
    notes = [
        "使い方：コントロールライン（最終周はゴール）を通った時の、スタートからの経過時間を「実走タイム（累計）」に入れる（例 0:04:57.1）。ほかは自動",
        "差：目標タイムより早ければ＋（緑）、遅れていれば−（赤）",
        "次の周の目標速度：次の周をこの速さで走ると、次の周の目標タイムちょうどに戻る（遅れ・貯金を次の1周で取り返す）",
        "次の周の前半の目標速度：遅れ・貯金を次の周の前半の半周で取り返すときの速さ。後半は計画の速さで走る",
        "一周平均速度の色：その周の目標速度以上〜＋1km/hで緑、目標未満で赤、それより速いと黄（速すぎて燃料がもったいない）",
        "周の距離：1周目はスタートからコントロールまで、最終周はコントロールからゴールまでなので短い。元の記録表は全周を同じ長さ（全体距離÷周回数）にしていた",
        "Androidアプリ（EcoLogger）の「この周の目標」「差」「前半」と同じ計算",
    ]
    for i, t in enumerate(notes):
        ws.cell(n + i, 2, t).font = NOTE

    widths = {"A": 2, "B": 10, "C": 11, "D": 13, "E": 13, "F": 10, "G": 11, "H": 11, "I": 11, "J": 10, "K": 24, "L": 13, "M": 12, "N": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C17"
    style_all(ws)
    return ws


race_sheet("CN二輪車", 3, 16, 27.36, 6856.72)
race_sheet("電動二輪", 3, 16, 27.36, 6856.72)
race_sheet("CNグループⅢ", 7, 39, 20.11, 16389.68)

# Last year's run, as the worked example and the reason for the change.
ws = wb.create_sheet("昨年の実走（検証）")
ws["B2"] = "昨年のスーパーカブ（2025-10-12 全国大会、GPXから）"; ws["B2"].font = TITLE
ws["B3"] = ("スタート・コントロール・ゴールを通った時刻をGPX（10〜16秒おきの点）から読み、1周ずつの平均速度を2通りの距離で出した。"
            "距離の補正後は3周ともほぼ同じ速さ＝実際は一定のペースで走れていた。等分のままだと1周目が速すぎ・2周目が遅すぎに見える")
ws["B3"].font = NOTE
head = ["周回", "実走タイム（累計）", "周タイム", "等分の距離（m）", "等分での平均速度", "補正後の距離（m）", "補正後の平均速度",
        "GPSの道のり（m）"]
for j, h in enumerate(head):
    c = ws.cell(5, 2 + j, h); c.font = BOLD; c.fill = HEAD; c.border = BOX
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
ws.row_dimensions[5].height = 32
times = [(0, 4, 57.1), (0, 10, 36.0), (0, 16, 0.9)]
gps = [2091, 2375, 2275]
for k in range(3):
    r = 6 + k
    ws.cell(r, 2, k + 1)
    h, m, s = times[k]
    c = ws.cell(r, 3, f"=TIME({h},{m},{int(s)})+{round(s - int(s), 1)}/86400"); c.font = BLUE; c.number_format = TIME_FMT
    ws.cell(r, 4, f"=C{r}-C{r - 1}" if k else f"=C{r}").number_format = TIME_FMT
    ws.cell(r, 5, "='CN二輪車'!$M$4/'CN二輪車'!$M$5").number_format = M_FMT
    ws.cell(r, 6, f"=E{r}/1000/(D{r}*24)").number_format = KMH_FMT
    ws.cell(r, 7, f"='CN二輪車'!C{18 + k}").number_format = M_FMT
    ws.cell(r, 8, f"=G{r}/1000/(D{r}*24)").number_format = KMH_FMT
    c = ws.cell(r, 9, gps[k]); c.font = BLUE; c.number_format = "#,##0"
    for j in range(2, 10):
        ws.cell(r, j).border = BOX
ws["B10"] = "ゴールは16:00.9、制限16:27.36に対して約26秒早い（余裕の既定30秒の根拠）"
ws["B11"] = "GPSの道のりは点が粗くカーブで角を切るため約2%短く出る。補正後の距離と2%以内で一致し、等分の距離（2285.6m）とは1周目で約190m違う"
ws["B12"] = "距離の補正後の値は「CN二輪車」シートの入力（スタート→コントロール・ゴール→コントロール）を使う"
for c in ("B10", "B11", "B12"):
    ws[c].font = NOTE
ws.cell(6, 9).comment = Comment("2025-10-12のGPX（iGPSPORT）で、各ラインを通った時刻の間の折れ線の長さ", "EcoLogger")
for col, w in {"A": 2, "B": 8, "C": 14, "D": 12, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13}.items():
    ws.column_dimensions[col].width = w
style_all(ws)

# Legend.
ws = wb.create_sheet("使い方", 0)
lines = [
    ("本番用記録表（改良版）", TITLE),
    ("元の記録表（Honda エコ マイレッジ チャレンジ　本番用記録表.xlsx）と同じく、実走タイムを入れると差と次の周の目標速度が出る。", BLACK),
    ("", BLACK),
    ("入れるセル", BOLD),
    ("・黄色のセル（青い文字）だけを書き換える。ほかのセルは計算式", BLACK),
    ("・実走タイム（累計）：コントロール（最終周はゴール）を通った時の、スタートからの経過時間。例 0:04:57.1", BLACK),
    ("・余裕（秒）：制限タイムよりどれだけ早くゴールするか。クラスのシートごとに決める", BLACK),
    ("・スタート→コントロール、ゴール→コントロール（m）：コースウォークで測ったラインの間隔に書き換える（アプリの設定→ラインの「ライン間の距離」）", BLACK),
    ("", BLACK),
    ("元の記録表から変えたところ", BOLD),
    ("1. 周の距離を実際のコースに合わせた：1周目はスタート→コントロール、最終周はコントロール→ゴールの分だけ短い（昨年の実走のGPXで確認、「昨年の実走（検証）」シート）", BLACK),
    ("2. 目標タイムを周回数で等分せず、距離に比例して割り振った（等分だと1周目に約20秒の「貯金」が出て、2周目を遅く走らせてしまう）", BLACK),
    ("3. 余裕を秒で入れる欄を作った（元は全体目標タイム16:20を直接入れていた）", BLACK),
    ("4. 一周平均速度の色を、決まった25〜26km/hではなく、その周の目標速度に対して付けた", BLACK),
    ("5. その周の目標速度と目標比の列を足した（ふり返り用）", BLACK),
    ("6. 元の記録表で直したもの：全体平均速度の式が周回数3を決め打ちしていた／説明文の25.45km/hが計算と合っていなかった（計画は25.19km/h）／"
     "最終周の行にGoogleスプレッドシート専用の関数（SPARKLINE）が入っていてExcelでは空になっていた", BLACK),
    ("7. CNグループⅢ（7周）のシートを足した", BLACK),
    ("", BLACK),
    ("前提", BOLD),
    ("・全体距離・周回数・制限タイム：2026年 全国大会 競技規則 第2章第1条第1項4", BLACK),
    ("・ライン間の距離の初期値：Googleマップで拾ったもてぎのおおよそのライン位置（2026-10-19）からの直線距離。コースウォークで測り直す", BLACK),
    ("・余裕の初期値30秒：昨年のスーパーカブが制限より約26秒早くゴールしていたことから", BLACK),
    ("・計算はAndroidアプリ（EcoLogger）の画面の「この周の目標」「差」「前半」と同じ", BLACK),
]
for i, (t, f) in enumerate(lines):
    c = ws.cell(2 + i, 2, t); c.font = f
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 150
style_all(ws)

wb.save(OUT)
print("saved", OUT)
